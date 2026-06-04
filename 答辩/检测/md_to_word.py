import re
import os
import sys
from docx import Document
from docx.shared import RGBColor, Pt

def parse_inline(text, is_blue=False):
    segments = []
    pattern = re.compile(r'\*\*(.*?)\*\*')
    last_end = 0
    for match in pattern.finditer(text):
        if match.start() > last_end:
            segments.append({'text': text[last_end:match.start()], 'blue': is_blue, 'bold': False})
        segments.append({'text': match.group(1), 'blue': is_blue, 'bold': True})
        last_end = match.end()
    if last_end < len(text):
        segments.append({'text': text[last_end:], 'blue': is_blue, 'bold': False})
    return segments

def parse_line(line):
    segments = []
    pattern = re.compile(r'<span style="color: blue;">(.*?)</span>', re.IGNORECASE)
    last_end = 0
    for match in pattern.finditer(line):
        if match.start() > last_end:
            segments.extend(parse_inline(line[last_end:match.start()], is_blue=False))
        segments.extend(parse_inline(match.group(1), is_blue=True))
        last_end = match.end()
    if last_end < len(line):
        segments.extend(parse_inline(line[last_end:], is_blue=False))
    return segments

def convert_md_to_docx(md_path, docx_path):
    doc = Document()
    
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
        
    # Split by double newline to get paragraphs
    blocks = content.split('\n\n')
    
    for block in blocks:
        block = block.strip()
        if not block:
            continue
            
        # Ignore images
        if block.startswith('![') or block.startswith('<img'):
            continue
            
        # Ignore tables (very basic check)
        if block.startswith('|') and '|' in block:
            continue
            
        # Join lines within a block (remove single newlines)
        block = block.replace('\n', '')
        
        # Strip blockquotes
        if block.startswith('> '):
            block = block[2:]
            
        # Check headings
        heading_match = re.match(r'^(#{1,6})\s+(.*)', block)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            
            # Headings can also have blue text or bold, we process it and add to heading
            h = doc.add_heading(level=level)
            segments = parse_line(text)
            for seg in segments:
                run = h.add_run(seg['text'])
                if seg['bold']:
                    run.bold = True
                if seg['blue']:
                    run.font.color.rgb = RGBColor(0, 0, 255)
            continue
            
        # Add a paragraph
        p = doc.add_paragraph()
        
        # Parse inline formatting
        segments = parse_line(block)
        for seg in segments:
            run = p.add_run(seg['text'])
            if seg['bold']:
                run.bold = True
            if seg['blue']:
                run.font.color.rgb = RGBColor(0, 0, 255)
                
    doc.save(docx_path)
    print(f"Successfully saved to {docx_path}")

if __name__ == '__main__':
    md_file = r'd:\office\jushuang1\github\ss\答辩\检测\v20p\73论文_v20p.md'
    docx_file = r'd:\office\jushuang1\github\ss\答辩\检测\v20p\73论文_v20p.docx'
    
    if os.path.exists(md_file):
        convert_md_to_docx(md_file, docx_file)
    else:
        print(f"Error: Could not find {md_file}")
