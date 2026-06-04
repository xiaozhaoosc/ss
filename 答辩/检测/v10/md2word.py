#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Markdown to Word 转换脚本
注意段落与换行的处理
"""

import os
import re
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
except ImportError:
    print("错误: 需要安装 python-docx 库")
    print("请运行: pip install python-docx")
    exit(1)

def parse_markdown(md_content):
    """
    解析Markdown内容，返回结构化的段落列表
    处理段落与换行
    """
    lines = md_content.split('\n')
    elements = []
    current_paragraph = []
    in_code_block = False
    code_block_content = []
    in_table = False
    table_rows = []

    for line in lines:
        # 处理代码块
        if line.strip().startswith('```'):
            if in_code_block:
                # 结束代码块
                elements.append({
                    'type': 'code',
                    'content': '\n'.join(code_block_content)
                })
                code_block_content = []
                in_code_block = False
            else:
                # 开始代码块
                in_code_block = True
            continue

        if in_code_block:
            code_block_content.append(line)
            continue

        # 处理表格
        if '|' in line and line.strip().startswith('|'):
            if not in_table:
                in_table = True
                table_rows = []
            # 跳过分隔行
            if re.match(r'^[\s|:-]+$', line):
                continue
            # 解析表格行
            row = [cell.strip() for cell in line.split('|')[1:-1]]
            table_rows.append(row)
            continue
        elif in_table:
            # 表格结束
            elements.append({
                'type': 'table',
                'content': table_rows
            })
            table_rows = []
            in_table = False

        # 处理空行（段落分隔）
        if not line.strip():
            if current_paragraph:
                elements.append({
                    'type': 'paragraph',
                    'content': ' '.join(current_paragraph)
                })
                current_paragraph = []
            continue

        # 处理标题
        if line.startswith('#'):
            level = len(line.split(' ')[0])
            title = line.lstrip('#').strip()
            elements.append({
                'type': 'heading',
                'level': level,
                'content': title
            })
            continue

        # 处理分隔线
        if line.strip() in ['---', '***', '___']:
            elements.append({
                'type': 'separator'
            })
            continue

        # 处理列表项
        if re.match(r'^[\s]*[-*+]\s', line) or re.match(r'^[\s]*\d+\.\s', line):
            if current_paragraph:
                elements.append({
                    'type': 'paragraph',
                    'content': ' '.join(current_paragraph)
                })
                current_paragraph = []
            elements.append({
                'type': 'list_item',
                'content': line.strip()
            })
            continue

        # 处理图片
        if re.match(r'^!\[.*\]\(.*\)$', line.strip()):
            elements.append({
                'type': 'image',
                'content': line.strip()
            })
            continue

        # 处理链接
        if re.match(r'^\[.*\]\(.*\)$', line.strip()):
            elements.append({
                'type': 'link',
                'content': line.strip()
            })
            continue

        # 普通段落
        current_paragraph.append(line.strip())

    # 处理最后一个段落
    if current_paragraph:
        elements.append({
            'type': 'paragraph',
            'content': ' '.join(current_paragraph)
        })

    # 处理未关闭的表格
    if in_table and table_rows:
        elements.append({
            'type': 'table',
            'content': table_rows
        })

    return elements

def create_word_document(elements, output_path):
    """
    创建Word文档
    处理段落与换行
    """
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    # 定义标题样式
    heading_styles = {
        1: ('Heading 1', Pt(22), True),
        2: ('Heading 2', Pt(16), True),
        3: ('Heading 3', Pt(14), True),
        4: ('Heading 4', Pt(13), True),
        5: ('Heading 5', Pt(12), True),
        6: ('Heading 6', Pt(11), True),
    }

    for element in elements:
        elem_type = element.get('type')

        if elem_type == 'heading':
            level = element.get('level', 1)
            content = element.get('content', '')

            # 限制标题级别在1-6之间
            level = max(1, min(6, level))

            # 创建标题段落
            heading = doc.add_heading(level=level)
            run = heading.add_run(content)
            run.bold = True

            # 设置标题字体大小
            if level in heading_styles:
                _, font_size, _ = heading_styles[level]
                run.font.size = font_size

        elif elem_type == 'paragraph':
            content = element.get('content', '')
            # 处理段落中的换行
            # 将 \n 替换为 Word 的换行符
            content = content.replace('\\n', '\n')

            # 创建段落
            para = doc.add_paragraph()

            # 处理段落中的多行文本
            lines = content.split('\n')
            for i, line in enumerate(lines):
                if line.strip():
                    run = para.add_run(line.strip())
                    run.font.size = Pt(12)

                # 在行之间添加换行（除了最后一行）
                if i < len(lines) - 1 and lines[i + 1].strip():
                    run.add_break()

        elif elem_type == 'code':
            content = element.get('content', '')
            # 创建代码块段落
            para = doc.add_paragraph()
            run = para.add_run(content)
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
            # 设置代码块背景色（灰色）
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            shading_elm = OxmlElement('w:shd')
            shading_elm.set(qn('w:fill'), 'F5F5F5')
            para.paragraph_format.element.get_or_add_pPr().append(shading_elm)

        elif elem_type == 'list_item':
            content = element.get('content', '')
            # 创建列表项
            para = doc.add_paragraph(style='List Bullet')
            # 移除列表项标记
            content = re.sub(r'^[-*+]\s', '', content)
            content = re.sub(r'^\d+\.\s', '', content)
            run = para.add_run(content.strip())
            run.font.size = Pt(12)

        elif elem_type == 'table':
            table_data = element.get('content', [])
            if table_data:
                # 创建表格
                rows = len(table_data)
                cols = len(table_data[0]) if table_data else 0
                if rows > 0 and cols > 0:
                    table = doc.add_table(rows=rows, cols=cols)
                    table.style = 'Table Grid'

                    # 填充表格内容
                    for i, row_data in enumerate(table_data):
                        for j, cell_content in enumerate(row_data):
                            if j < cols:
                                cell = table.cell(i, j)
                                cell.text = cell_content

                                # 设置表格字体
                                for paragraph in cell.paragraphs:
                                    for run in paragraph.runs:
                                        run.font.size = Pt(10)

        elif elem_type == 'separator':
            # 添加分隔线（使用水平线）
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(6)
            para.paragraph_format.space_before = Pt(6)
            # 添加水平线
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), 'auto')
            pBdr.append(bottom)
            para.paragraph_format.element.get_or_add_pPr().append(pBdr)

    # 保存文档
    doc.save(output_path)
    return True

def main():
    """
    主函数
    """
    print("=" * 60)
    print("  Markdown to Word 转换工具 (Python版)")
    print("=" * 60)
    print()

    # 获取脚本所在目录
    script_dir = Path(__file__).parent

    # 输入和输出文件
    input_file = script_dir / "论文_v10.md"
    output_file = script_dir / "论文_v10.docx"

    # 检查输入文件是否存在
    if not input_file.exists():
        print(f"[错误] 找不到输入文件: {input_file}")
        return False

    print(f"[信息] 输入文件: {input_file.name}")
    print(f"[信息] 输出文件: {output_file.name}")
    print()

    # 读取Markdown内容
    print("[执行] 读取Markdown文件...")
    try:
        with open(input_file, 'r', encoding='utf-8') as f:
            md_content = f.read()
    except Exception as e:
        print(f"[错误] 读取文件失败: {e}")
        return False

    # 解析Markdown
    print("[执行] 解析Markdown内容...")
    elements = parse_markdown(md_content)
    print(f"[信息] 解析完成，共 {len(elements)} 个元素")

    # 创建Word文档
    print("[执行] 创建Word文档...")
    try:
        success = create_word_document(elements, output_file)
        if success:
            print()
            print("[成功] 转换完成！")
            print(f"[输出] {output_file}")

            # 显示文件信息
            file_size = output_file.stat().st_size
            file_size_kb = round(file_size / 1024, 2)
            print(f"[信息] 文件大小: {file_size_kb} KB")

            return True
        else:
            print("[错误] 转换失败")
            return False
    except Exception as e:
        print(f"[错误] 创建Word文档失败: {e}")
        return False

if __name__ == "__main__":
    success = main()
    if not success:
        exit(1)
