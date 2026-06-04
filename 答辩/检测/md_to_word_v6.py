import re, os
from docx import Document
from docx.shared import RGBColor

# ---------- 辅助解析函数 ----------
def parse_inline(text, blue=False):
    segs = []
    # 加粗 **text**
    bold_pat = re.compile(r'\*\*(.*?)\*\*')
    pos = 0
    for m in bold_pat.finditer(text):
        if m.start() > pos:
            segs.append({'txt': text[pos:m.start()], 'blue': blue, 'bold': False})
        segs.append({'txt': m.group(1), 'blue': blue, 'bold': True})
        pos = m.end()
    if pos < len(text):
        segs.append({'txt': text[pos:], 'blue': blue, 'bold': False})
    return segs

def parse_line(line):
    segs = []
    blue_pat = re.compile(r'<span style="color: blue;">(.*?)</span>', re.IGNORECASE)
    pos = 0
    for m in blue_pat.finditer(line):
        if m.start() > pos:
            segs.extend(parse_inline(line[pos:m.start()], blue=False))
        segs.extend(parse_inline(m.group(1), blue=True))
        pos = m.end()
    if pos < len(line):
        segs.extend(parse_inline(line[pos:], blue=False))
    return segs

def md_to_docx(md_path, docx_path):
    doc = Document()
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
    # 按双换行划分块，形成段落
    blocks = content.split('\n\n')
    for blk in blocks:
        blk = blk.strip()
        if not blk:
            continue
        # 忽略图片、表格
        if blk.startswith('![') or blk.startswith('|'):
            continue
        # 标题处理
        h_match = re.match(r'^(#{1,6})\s+(.*)', blk)
        if h_match:
            level = len(h_match.group(1))
            txt = h_match.group(2)
            heading = doc.add_heading(level=level)
            for seg in parse_line(txt):
                run = heading.add_run(seg['txt'])
                if seg['bold']:
                    run.bold = True
                if seg['blue']:
                    run.font.color.rgb = RGBColor(0,0,255)
            continue
        # 普通段落
        p = doc.add_paragraph()
        for seg in parse_line(blk):
            run = p.add_run(seg['txt'])
            if seg['bold']:
                run.bold = True
            if seg['blue']:
                run.font.color.rgb = RGBColor(0,0,255)
    doc.save(docx_path)
    print(f'Saved → {docx_path}')

if __name__ == '__main__':
    md_file = r'd:\office\jushuang1\github\ss\答辩\检测\v6\73论文_v6.md'
    docx_file = r'd:\office\jushuang1\github\ss\答辩\检测\v6\73论文_v6.docx'
    md_to_docx(md_file, docx_file)
