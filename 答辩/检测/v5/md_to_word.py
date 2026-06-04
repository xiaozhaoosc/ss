#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Markdown 转 Word 脚本
功能：
1. 读取 v5_论文.md 文件
2. 转换为 .docx 格式
3. 忽略表格和图片
4. 注意段落与换行的处理

依赖：pip install python-docx
"""

import re
import os
import sys

try:
    from docx import Document
    from docx.shared import Pt, Cm, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
except ImportError:
    print('需要安装 python-docx 库')
    print('正在安装...')
    os.system(f'{sys.executable} -m pip install python-docx -q')
    from docx import Document
    from docx.shared import Pt, Cm, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn


def set_font(run, name_cn='宋体', name_en='Times New Roman', size=12, bold=False):
    """设置字体"""
    run.font.name = name_en
    run.font.size = Pt(size)
    run.bold = bold
    # 设置中文字体
    r = run._element
    rPr = r.find(qn('w:rPr'))
    if rPr is None:
        rPr = r.makeelement(qn('w:rPr'), {})
        r.insert(0, rPr)
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), name_cn)


def set_paragraph_format(para, first_line_indent=None, line_spacing=1.5,
                          space_before=0, space_after=0, alignment=None):
    """设置段落格式"""
    fmt = para.paragraph_format
    fmt.line_spacing = line_spacing
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    if first_line_indent is not None:
        fmt.first_line_indent = Cm(first_line_indent)
    if alignment is not None:
        fmt.alignment = alignment


def is_heading(line):
    """判断是否为标题行"""
    # 匹配 "1绪论", "1.1研究背景", "（1）xxx" 等格式
    patterns = [
        r'^(\d+)\s*[^\d\.\s]',         # "1绪论", "2系统分析"
        r'^(\d+\.\d+)\s*[^\d\.\s]',     # "1.1研究背景"
        r'^(\d+\.\d+\.\d+)\s*[^\d\.\s]', # "1.4.1传统行为干预"
        r'^（\d+）',                     # "（1）xxx"
    ]
    stripped = line.strip().lstrip('\u3000').strip()
    for p in patterns:
        if re.match(p, stripped):
            return True
    return False


def get_heading_level(line):
    """获取标题级别"""
    stripped = line.strip().lstrip('\u3000').strip()
    if re.match(r'^(\d+\.\d+\.\d+)', stripped):
        return 3
    elif re.match(r'^(\d+\.\d+)', stripped):
        return 2
    elif re.match(r'^\d+\s*[^\d\.\s]', stripped):
        return 1
    elif re.match(r'^（\d+）', stripped):
        return 4
    return 0


def is_table_line(line):
    """判断是否为表格行"""
    stripped = line.strip()
    # 表格标题行
    if re.match(r'^表\s*\d+', stripped):
        return True
    # 表格数据行（包含多个空格分隔的列）
    if re.match(r'^\s*字段名称\s+类型', stripped):
        return True
    # 以字段名开头的数据行
    if re.match(r'^\s*\w+\s+(int[248]|varchar|char|timestamp|date|jsonb|text|bool)', stripped):
        return True
    return False


def is_image_line(line):
    """判断是否为图片行"""
    stripped = line.strip()
    # "图 X.X xxx" 格式
    if re.match(r'^图\s*\d+', stripped):
        return True
    # 空行（可能是图片占位）
    if stripped == '':
        return True
    return False


def is_skip_line(line):
    """判断是否应跳过的行"""
    stripped = line.strip()
    # 表格相关
    if is_table_line(line):
        return True
    # 纯图片引用行
    if stripped.startswith('图 ') and len(stripped) < 30:
        return True
    return False


def is_special_section(line):
    """判断是否为特殊段落（封面、目录等）"""
    stripped = line.strip()
    special = [
        '上海应用技术大学',
        'SHANGHAI INSTITUTE OF TECHNOLOGY',
        '高等学历继续教育',
        '本科毕业设计（论文）',
        '课题名称',
        '专    业',
        '班    级',
        '学生学号',
        '学生姓名',
        '指导教师',
        '2026年',
        'ADHD中小学生（儿童）行为习惯辅助系统设计与实现',
        '目  录',
    ]
    for s in special:
        if stripped.startswith(s):
            return True
    return False


def md_to_docx(md_path, docx_path):
    """将 Markdown 文件转换为 Word 文档"""
    print(f'读取: {md_path}')

    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    # 清理所有HTML span标签
    content = re.sub(r'<span[^>]*>', '', content)
    content = re.sub(r'</span>', '', content)

    lines = content.split('\n')
    print(f'总行数: {len(lines)}')

    doc = Document()

    # 设置默认样式
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    style.paragraph_format.line_spacing = 1.5

    # 设置页边距
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)

    in_table_block = False
    skip_count = 0
    para_count = 0

    for i, line in enumerate(lines):
        raw_line = line.rstrip('\r')

        # 跳过空行
        if raw_line.strip() == '':
            continue

        # 跳过纯图片行
        stripped = raw_line.strip().lstrip('\u3000').strip()
        if re.match(r'^图\s*\d+[\.\d]*\s*\S', stripped) and len(stripped) < 50:
            skip_count += 1
            continue

        # 跳过表格数据行
        if is_table_line(raw_line):
            in_table_block = True
            skip_count += 1
            continue

        # 表格块结束检测
        if in_table_block:
            if not is_table_line(raw_line) and stripped:
                in_table_block = False
            else:
                skip_count += 1
                continue

        # 表格标题行（如 "表 3.1系统用户表"）
        if re.match(r'^表\s*\d+', stripped):
            skip_count += 1
            continue

        # 处理特殊段落（封面信息）
        if is_special_section(raw_line):
            para = doc.add_paragraph()
            if '上海应用技术大学' in stripped or 'SHANGHAI' in stripped:
                set_paragraph_format(para, alignment=WD_ALIGN_PARAGRAPH.CENTER)
                run = para.add_run(stripped)
                set_font(run, size=16, bold=True)
            elif stripped.startswith('ADHD中小学生'):
                set_paragraph_format(para, alignment=WD_ALIGN_PARAGRAPH.CENTER)
                run = para.add_run(stripped)
                set_font(run, name_cn='黑体', size=16, bold=True)
            else:
                set_paragraph_format(para, alignment=WD_ALIGN_PARAGRAPH.CENTER)
                run = para.add_run(stripped)
                set_font(run, size=14)
            para_count += 1
            continue

        # 处理目录行（长文本，包含大量章节号）
        if stripped.startswith('目  录') or (len(stripped) > 100 and re.search(r'\d+\.\d+', stripped) and '绪论' in stripped):
            # 跳过目录内容
            skip_count += 1
            continue

        # 处理标题
        if is_heading(raw_line):
            level = get_heading_level(raw_line)
            para = doc.add_paragraph()

            if level == 1:
                set_paragraph_format(para, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                                      space_before=12, space_after=6)
                run = para.add_run(stripped)
                set_font(run, name_cn='黑体', size=16, bold=True)
            elif level == 2:
                set_paragraph_format(para, space_before=6, space_after=3)
                run = para.add_run(stripped)
                set_font(run, name_cn='黑体', size=14, bold=True)
            elif level == 3:
                set_paragraph_format(para, space_before=3, space_after=3)
                run = para.add_run(stripped)
                set_font(run, name_cn='黑体', size=13, bold=True)
            else:
                set_paragraph_format(para, first_line_indent=0.74,
                                      space_before=3, space_after=3)
                run = para.add_run(stripped)
                set_font(run, name_cn='黑体', size=12, bold=True)

            para_count += 1
            continue

        # 处理关键词行
        if stripped.startswith('关键词') or stripped.startswith('Keywords'):
            para = doc.add_paragraph()
            set_paragraph_format(para, first_line_indent=0.74)
            run = para.add_run(stripped)
            set_font(run, size=12, bold=True)
            para_count += 1
            continue

        # 处理摘要标识
        if stripped.startswith('摘要') or stripped.startswith('Abstract'):
            para = doc.add_paragraph()
            # 摘要带首行缩进
            set_paragraph_format(para, first_line_indent=0.74)
            # 分离"摘要："标签和正文
            if '：' in stripped or ':' in stripped:
                sep = '：' if '：' in stripped else ':'
                label, body = stripped.split(sep, 1)
                run_label = para.add_run(label + sep)
                set_font(run_label, size=12, bold=True)
                if body.strip():
                    # 去掉可能的tab
                    body = body.replace('\t', '')
                    run_body = para.add_run(body.strip())
                    set_font(run_body, size=12)
            else:
                run = para.add_run(stripped)
                set_font(run, size=12)
            para_count += 1
            continue

        # 处理致谢、参考文献、附录等特殊标题
        if stripped in ['致谢', '参考文献', '附  录', '附录']:
            para = doc.add_paragraph()
            set_paragraph_format(para, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                                  space_before=12, space_after=6)
            run = para.add_run(stripped)
            set_font(run, name_cn='黑体', size=16, bold=True)
            para_count += 1
            continue

        # 普通段落
        # 处理全角空格开头的缩进
        text = raw_line
        # 移除开头的全角空格（用于首行缩进）
        text = text.lstrip('\u3000').strip()

        if text:
            para = doc.add_paragraph()
            set_paragraph_format(para, first_line_indent=0.74)
            run = para.add_run(text)
            set_font(run, size=12)
            para_count += 1

    # 保存
    print(f'\n段落数: {para_count}')
    print(f'跳过行数: {skip_count}')
    print(f'保存到: {docx_path}')
    doc.save(docx_path)
    print('完成!')


def main():
    md_path = r'd:\office\jushuang1\github\ss\答辩\检测\v5\v5_论文.md'
    docx_path = r'd:\office\jushuang1\github\ss\答辩\检测\v5\v5_论文.docx'

    if not os.path.exists(md_path):
        print(f'错误: 找不到文件 {md_path}')
        print('请先运行 rewrite_v5.py 生成 v5_论文.md')
        sys.exit(1)

    md_to_docx(md_path, docx_path)


if __name__ == '__main__':
    main()
