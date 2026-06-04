#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
v4.md -> v4.docx 转换脚本
处理要点：
1. 段落与换行：MD中的连续段落正确转换为Word段落
2. 标题层级：# ## ### #### 映射到Heading 1-4
3. 代码块：保留等宽字体格式
4. 表格：转换为Word表格
5. 列表：转换为有序/无序列表
6. 优先使用pandoc，失败则用python-docx回退
"""

import os
import re
import sys
import subprocess
import io

# 解决Windows控制台中文乱码
if sys.platform == 'win32':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')


def log(msg):
    print(msg, flush=True)


def convert_with_pandoc(md_path, docx_path):
    """使用pandoc转换"""
    ref_doc = os.path.join(os.path.dirname(md_path), "reference.docx")
    if os.path.exists(ref_doc):
        cmd = ["pandoc", "-f", "markdown+smart", "-t", "docx",
               "--reference-doc", ref_doc, "-o", docx_path, md_path]
    else:
        cmd = ["pandoc", "-f", "markdown+smart", "-t", "docx",
               "-o", docx_path, md_path]

    log("executing: " + " ".join(cmd))
    result = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8', errors='replace')

    if result.returncode == 0:
        log("[OK] pandoc conversion succeeded: " + docx_path)
        return True
    else:
        log("[FAIL] pandoc error: " + result.stderr)
        return False


def convert_with_python_docx(md_path, docx_path):
    """使用python-docx转换（回退方案）"""
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        log("installing python-docx...")
        subprocess.run([sys.executable, "-m", "pip", "install", "python-docx", "--quiet"])
        from docx import Document
        from docx.shared import Pt

    doc = Document()

    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    content = content.replace('\r\n', '\n').replace('\r', '\n')
    lines = content.split('\n')

    i = 0
    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal table_rows
        if not table_rows:
            return
        num_cols = max(len(r) for r in table_rows)
        tbl = doc.add_table(rows=len(table_rows), cols=num_cols)
        tbl.style = 'Table Grid'
        for ri, row_data in enumerate(table_rows):
            row = tbl.rows[ri]
            for ci, cell_text in enumerate(row_data):
                if ci < num_cols:
                    cell = row.cells[ci]
                    cell.text = cell_text
                    if ri == 0:
                        for para in cell.paragraphs:
                            for run in para.runs:
                                run.bold = True
        table_rows = []

    def add_inline_runs(paragraph, text):
        """处理行内粗体/斜体"""
        parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('*') and part.endswith('*'):
                run = paragraph.add_run(part[1:-1])
                run.italic = True
            else:
                if part:
                    paragraph.add_run(part)

    def clean_md(text):
        """移除HTML标签"""
        text = re.sub(r'<[^>]+>', '', text)
        return text

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # --- 代码块 ---
        if stripped.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lines = []
                i += 1
                continue
            else:
                in_code_block = False
                if code_lines:
                    p = doc.add_paragraph()
                    run = p.add_run('\n'.join(code_lines))
                    run.font.name = 'Courier New'
                    run.font.size = Pt(9)
                code_lines = []
                i += 1
                continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # --- 表格行 ---
        if stripped.startswith('|') and stripped.endswith('|'):
            cells = [c.strip() for c in stripped.strip('|').split('|')]
            # 分隔行（---|---）跳过
            if all(re.match(r'^[-:]+$', c) for c in cells if c):
                i += 1
                continue
            table_rows.append(cells)
            i += 1
            continue
        else:
            flush_table()

        # --- 空行 ---
        if not stripped:
            i += 1
            continue

        # --- 标题 ---
        if stripped.startswith('#### '):
            doc.add_heading(stripped[5:], level=4)
        elif stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith('## '):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith('# '):
            doc.add_heading(stripped[2:], level=1)

        # --- 有序列表 ---
        elif re.match(r'^\d+\.\s', stripped):
            text = re.sub(r'^\d+\.\s', '', stripped)
            p = doc.add_paragraph(style='List Number')
            add_inline_runs(p, clean_md(text))

        # --- 无序列表 ---
        elif stripped.startswith('- ') or stripped.startswith('* '):
            text = stripped[2:]
            p = doc.add_paragraph(style='List Bullet')
            add_inline_runs(p, clean_md(text))

        # --- 分割线 ---
        elif stripped in ('---', '***', '___'):
            doc.add_paragraph('_' * 50)

        # --- 普通段落（合并后续连续非空行）---
        else:
            para_lines = [stripped]
            j = i + 1
            while j < len(lines):
                ns = lines[j].strip()
                if not ns:
                    break
                if (ns.startswith('#') or ns.startswith('|') or
                        ns.startswith('```') or re.match(r'^\d+\.\s', ns) or
                        ns.startswith('- ') or ns.startswith('* ')):
                    break
                para_lines.append(ns)
                j += 1
            full_text = ' '.join(para_lines)
            p = doc.add_paragraph()
            add_inline_runs(p, clean_md(full_text))
            i = j - 1

        i += 1

    flush_table()
    doc.save(docx_path)
    log("[OK] python-docx conversion succeeded: " + docx_path)
    return True


def main():
    current_dir = os.path.dirname(os.path.abspath(__file__))
    md_path = os.path.join(current_dir, "v4.md")
    docx_path = os.path.join(current_dir, "v4.docx")

    if not os.path.exists(md_path):
        log("ERROR: file not found: " + md_path)
        sys.exit(1)

    log("converting: " + md_path)
    log("output:     " + docx_path)
    log("-" * 50)

    if not convert_with_pandoc(md_path, docx_path):
        log("falling back to python-docx...")
        convert_with_python_docx(md_path, docx_path)

    if os.path.exists(docx_path):
        size_kb = os.path.getsize(docx_path) / 1024
        log(f"conversion complete! size: {size_kb:.1f} KB")
    else:
        log("ERROR: output file not found")
        sys.exit(1)


if __name__ == "__main__":
    main()
