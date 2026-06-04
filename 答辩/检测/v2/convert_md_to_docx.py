import os
import subprocess
import sys

def convert():
    current_dir = os.path.dirname(os.path.abspath(__file__))
    md_path = os.path.join(current_dir, "v2.md")
    docx_path = os.path.join(current_dir, "v2.docx")
    
    if not os.path.exists(md_path):
        print(f"Error: {md_path} not found!")
        sys.exit(1)
        
    print(f"Starting conversion: {md_path} -> {docx_path}")
    
    # 优先使用系统已有的 pandoc 进行高质量转换
    try:
        cmd = ["pandoc", "-f", "markdown", "-t", "docx", md_path, "-o", docx_path]
        print(f"Executing: {' '.join(cmd)}")
        result = subprocess.run(cmd, capture_output=True, text=True)
        if result.returncode == 0:
            print("Conversion completed successfully via Pandoc!")
            print(f"Output saved to: {docx_path}")
            return
        else:
            print(f"Pandoc error code: {result.returncode}")
            print(f"Stderr: {result.stderr}")
    except Exception as e:
        print(f"Pandoc call failed: {e}")
        
    # 如果 pandoc 失败或不可用，作为首席开发代理，提供鲁棒的 python-docx 后备机制
    print("Trying backup conversion method via python-docx...")
    try:
        import docx
    except ImportError:
        print("python-docx not installed, attempting auto-installation...")
        subprocess.run([sys.executable, "-m", "pip", "install", "python-docx", "--quiet"])
        import docx
        
    try:
        doc = docx.Document()
        with open(md_path, "r", encoding="utf-8") as f:
            lines = f.readlines()
            
        print("Parsing Markdown and building Word Document...")
        for line in lines:
            stripped = line.strip()
            if not stripped:
                doc.add_paragraph("")
                continue
                
            # 处理标题
            if stripped.startswith("# "):
                doc.add_heading(stripped[2:], level=1)
            elif stripped.startswith("## "):
                doc.add_heading(stripped[3:], level=2)
            elif stripped.startswith("### "):
                doc.add_heading(stripped[4:], level=3)
            elif stripped.startswith("#### "):
                doc.add_heading(stripped[5:], level=4)
            else:
                # 简单文本段落，自动清除其中的任何 HTML span 标签以防乱入
                clean_line = stripped
                # 清理 HTML 标签
                import re
                clean_line = re.sub(r'<span.*?>', '', clean_line)
                clean_line = re.sub(r'</span>', '', clean_line)
                clean_line = re.sub(r'<font.*?>', '', clean_line)
                clean_line = re.sub(r'</font>', '', clean_line)
                
                doc.add_paragraph(clean_line)
                
        doc.save(docx_path)
        print("Backup conversion via python-docx completed successfully!")
        print(f"Output saved to: {docx_path}")
    except Exception as ex:
        print(f"Backup conversion failed: {ex}")
        sys.exit(1)

if __name__ == "__main__":
    convert()
