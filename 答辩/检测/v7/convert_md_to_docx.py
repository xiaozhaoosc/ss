from __future__ import annotations

from pathlib import Path
import argparse
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_MD = BASE_DIR / "论文_v7.md"
DEFAULT_DOCX = BASE_DIR / "论文_v7.docx"


HEADING1_RE = re.compile(r"^[1-6][^\d\.].{0,30}$")
HEADING2_RE = re.compile(r"^\d+\.\d+[^\d].*")
HEADING3_RE = re.compile(r"^\d+\.\d+\.\d+.*")
BLUE_OPEN_RE = re.compile(r'<span\s+style="color:\s*blue;">')


def set_font(run, east_asia="宋体", ascii_font="Times New Roman", size=12, bold=False):
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.bold = bold


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    for name, size in [("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.5


def is_cjk_or_punct(ch: str) -> bool:
    return bool(re.match(r"[\u4e00-\u9fff，。；：、！？）》】」』…—]", ch))


def join_lines(left: str, right: str) -> str:
    left = left.rstrip()
    right = right.lstrip()
    if not left:
        return right
    if not right:
        return left
    if is_cjk_or_punct(left[-1]) or is_cjk_or_punct(right[0]):
        return left + right
    return left + " " + right


def line_kind(line: str) -> str:
    stripped = line.strip()
    if not stripped:
        return "blank"
    if stripped.startswith("#"):
        return "heading"
    if HEADING3_RE.match(stripped):
        return "heading"
    if HEADING2_RE.match(stripped):
        return "heading"
    if HEADING1_RE.match(stripped):
        return "heading"
    if stripped in {"致谢", "参考文献", "附  录", "附录"}:
        return "heading"
    if stripped.startswith(("表 ", "图 ", "表", "图")) and len(stripped) <= 80:
        return "caption"
    if re.match(r"^(\d+\.|[（(]\d+[）)]|[一二三四五六七八九十]+、)", stripped):
        return "list"
    return "paragraph"


def read_blocks(md_path: Path) -> list[tuple[str, str]]:
    raw = md_path.read_text(encoding="utf-8").replace("\r\n", "\n").replace("\r", "\n")
    lines = raw.split("\n")
    blocks: list[tuple[str, str]] = []
    current = ""
    current_kind = "paragraph"

    def flush() -> None:
        nonlocal current, current_kind
        if current.strip():
            blocks.append((current_kind, current.strip()))
        current = ""
        current_kind = "paragraph"

    in_code = False
    for line in lines:
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            in_code = not in_code
            blocks.append(("code", stripped))
            continue
        if in_code:
            blocks.append(("code", line.rstrip()))
            continue

        kind = line_kind(line)
        if kind == "blank":
            flush()
            continue
        if kind in {"heading", "caption", "list"}:
            flush()
            blocks.append((kind, stripped))
            continue

        if not current:
            current = stripped
            current_kind = "paragraph"
        else:
            current = join_lines(current, stripped)

    flush()
    return blocks


def strip_markdown_heading(text: str) -> tuple[int, str]:
    if text.startswith("### "):
        return 3, text[4:].strip()
    if text.startswith("## "):
        return 2, text[3:].strip()
    if text.startswith("# "):
        return 1, text[2:].strip()
    if HEADING3_RE.match(text):
        return 3, text
    if HEADING2_RE.match(text):
        return 2, text
    return 1, text


def add_runs_with_blue(p, text: str, *, size=12, bold=False) -> None:
    parts = re.split(r'(<span\s+style="color:\s*blue;">|</span>)', text)
    blue = False
    for part in parts:
        if not part:
            continue
        if BLUE_OPEN_RE.fullmatch(part):
            blue = True
            continue
        if part == "</span>":
            blue = False
            continue
        clean = re.sub(r"<[^>]+>", "", part)
        if not clean:
            continue
        run = p.add_run(clean)
        set_font(run, size=size, bold=bold)
        if blue:
            run.font.color.rgb = RGBColor(0, 0, 255)


def add_body_paragraph(doc: Document, text: str, *, first_line_indent=True) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if first_line_indent:
        p.paragraph_format.first_line_indent = Pt(24)
    add_runs_with_blue(p, text)


def convert(md_path: Path, docx_path: Path) -> None:
    doc = Document()
    configure_styles(doc)
    blocks = read_blocks(md_path)

    before_abstract = True
    for kind, text in blocks:
        if text.startswith("摘要"):
            before_abstract = False

        if kind == "heading":
            level, title = strip_markdown_heading(text)
            p = doc.add_heading(level=min(level, 3))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
            add_runs_with_blue(p, title, size=16 if level == 1 else 14 if level == 2 else 12, bold=True)
            continue

        if before_abstract and len(text) <= 50:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs_with_blue(p, text, size=14 if "大学" in text or "设计" in text else 12, bold=False)
            continue

        if kind == "caption":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs_with_blue(p, text, size=10)
            continue

        if kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            add_runs_with_blue(p, text, size=10)
            continue

        if text.startswith(("关键词", "Keywords")) or text in {"致谢", "参考文献"}:
            add_body_paragraph(doc, text, first_line_indent=False)
        else:
            add_body_paragraph(doc, text, first_line_indent=True)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert v7 markdown thesis to Word docx.")
    parser.add_argument("md", nargs="?", default=str(DEFAULT_MD), help="Markdown input path")
    parser.add_argument("-o", "--output", default=str(DEFAULT_DOCX), help="DOCX output path")
    args = parser.parse_args()
    md_path = Path(args.md).resolve()
    docx_path = Path(args.output).resolve()
    if not md_path.exists():
        raise FileNotFoundError(f"Markdown file not found: {md_path}")
    convert(md_path, docx_path)
    print(f"Converted: {md_path} -> {docx_path}")


if __name__ == "__main__":
    main()
