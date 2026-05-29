from pptx import Presentation

prs = Presentation(r'd:\office\jushuang1\github\ss\答辩\WPS-ADHD儿童行为辅助系统答辩PPT_备注版.pptx')

with open(r'd:\office\jushuang1\github\ss\答辩\ppt_notes.txt', 'w', encoding='utf-8') as f:
    for i, slide in enumerate(prs.slides):
        f.write(f'=== Slide {i+1} ===\n')
        f.write('--- Slide Text ---\n')
        slide_text = '\n'.join([shape.text for shape in slide.shapes if shape.has_text_frame])
        f.write(slide_text + '\n')
        f.write('--- Notes ---\n')
        notes = slide.notes_slide.notes_text_repr.text if slide.has_notes_slide else '(no notes)'
        f.write(notes + '\n\n')

print('Done')